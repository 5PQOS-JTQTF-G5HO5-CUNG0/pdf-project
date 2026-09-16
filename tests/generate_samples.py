import io
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import fitz
import docx
from docx.shared import Inches, Pt, RGBColor
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor

SAMPLES_DIR = Path(__file__).parent / "samples"
SAMPLES_DIR.mkdir(parents=True, exist_ok=True)

def generate_sample_images():
    # 1. image.jpg
    img_jpg = Image.new("RGB", (800, 600), color=(240, 244, 248))
    draw = ImageDraw.Draw(img_jpg)
    draw.rectangle([50, 50, 750, 550], outline=(79, 70, 229), width=4)
    draw.text((100, 280), "LocalPDF Sample JPG Image (中文测试)", fill=(30, 41, 59))
    img_jpg.save(SAMPLES_DIR / "image.jpg", "JPEG", quality=95)

    # 2. image.png (带透明度或彩色图层)
    img_png = Image.new("RGBA", (600, 600), color=(255, 255, 255, 0))
    draw_png = ImageDraw.Draw(img_png)
    draw_png.ellipse([100, 100, 500, 500], fill=(99, 102, 241, 220), outline=(67, 56, 202, 255), width=6)
    draw_png.text((200, 290), "LocalPDF PNG", fill=(255, 255, 255, 255))
    img_png.save(SAMPLES_DIR / "image.png", "PNG")
    print("✓ 已生成 image.jpg 与 image.png")

def generate_docx_samples():
    # 1. simple.docx
    doc_simple = docx.Document()
    doc_simple.add_heading("LocalPDF 简单示例文档", 0)
    p = doc_simple.add_paragraph("这是一份用于验证 DOCX 转 PDF 的基础测试文档。")
    p.add_run("包含中文加粗字体").bold = True
    doc_simple.save(SAMPLES_DIR / "simple.docx")

    # 2. complex.docx
    doc_complex = docx.Document()
    doc_complex.add_heading("LocalPDF 复杂排版与报表测试", level=0)
    p1 = doc_complex.add_paragraph("本文件用于验证复杂表格、多段落与图片样式的转换质量。")

    # 添加表格
    table = doc_complex.add_table(rows=4, cols=4)
    table.style = 'Table Grid'
    headers = ["项目序号", "服务名称", "规格说明", "计费单价 (元)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
    rows_data = [
        ["001", "本地 Docker 容器转换引擎", "单节点离线部署", "0.00"],
        ["002", "Office 格式高保真转 PDF", "Gotenberg & LibreOffice 8", "0.00"],
        ["003", "AES-256 PDF 加解密与水印", "PyMuPDF 向量化渲染", "0.00"]
    ]
    for r_idx, r_data in enumerate(rows_data, start=1):
        for c_idx, val in enumerate(r_data):
            table.cell(r_idx, c_idx).text = val

    # 插入已生成的图片
    img_path = SAMPLES_DIR / "image.jpg"
    if img_path.exists():
        doc_complex.add_paragraph("下方为内嵌图片测试：")
        doc_complex.add_picture(str(img_path), width=Inches(4))

    doc_complex.save(SAMPLES_DIR / "complex.docx")
    print("✓ 已生成 simple.docx 与 complex.docx")

def generate_xlsx_samples():
    # 1. table.xlsx
    wb_simple = openpyxl.Workbook()
    ws = wb_simple.active
    ws.title = "简单列表"
    ws.append(["ID", "姓名", "部门", "入职日期"])
    ws.append([101, "张三", "技术研发部", "2024-01-15"])
    ws.append([102, "李四", "产品设计部", "2024-03-20"])
    wb_simple.save(SAMPLES_DIR / "table.xlsx")

    # 2. complex.xlsx
    wb_complex = openpyxl.Workbook()
    ws1 = wb_complex.active
    ws1.title = "月度财务报表"
    ws1.merge_cells("A1:E1")
    title_cell = ws1["A1"]
    title_cell.value = "2024年度本地财务结算汇总表"
    title_cell.font = Font(size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="4F46E5")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    headers = ["科目编号", "费用分类", "一季度预算", "实际支出", "预算执行率"]
    ws1.append([]) # 空行
    ws1.append(headers)
    header_fill = PatternFill("solid", fgColor="E0E7FF")
    for col in range(1, 6):
        cell = ws1.cell(row=3, column=col)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    rows = [
        ["6001", "服务器硬件折旧", 50000, 48500, "97%"],
        ["6002", "办公网络维护", 12000, 11800, "98.3%"],
        ["6003", "安全合规审查", 30000, 29000, "96.7%"]
    ]
    for r in rows:
        ws1.append(r)

    wb_complex.save(SAMPLES_DIR / "complex.xlsx")
    print("✓ 已生成 table.xlsx 与 complex.xlsx")

def generate_pptx_samples():
    # 1. simple.pptx
    prs_simple = Presentation()
    slide_layout = prs_simple.slide_layouts[0] # 标题页
    slide = prs_simple.slides.add_slide(slide_layout)
    slide.shapes.title.text = "LocalPDF 演示文稿"
    slide.placeholders[1].text = "简单 PPTX 转 PDF 格式验证"
    prs_simple.save(SAMPLES_DIR / "simple.pptx")

    # 2. complex.pptx
    prs_complex = Presentation()
    slide1 = prs_complex.slides.add_slide(prs_complex.slide_layouts[0])
    slide1.shapes.title.text = "LocalPDF 架构与核心优势"
    slide1.placeholders[1].text = "本地 Docker 隔离 · 隐私安全 · 全功能工具箱"

    slide2 = prs_complex.slides.add_slide(prs_complex.slide_layouts[1])
    slide2.shapes.title.text = "核心功能清单"
    tf = slide2.shapes.placeholders[1].text_frame
    tf.text = "1. Office 与 PDF 双向转换"
    p2 = tf.add_paragraph()
    p2.text = "2. PDF 高性能多页提取、拆分与旋转"
    p3 = tf.add_paragraph()
    p3.text = "3. AES-256 级别密码保护与文本防伪水印"

    prs_complex.save(SAMPLES_DIR / "complex.pptx")
    print("✓ 已生成 simple.pptx 与 complex.pptx")

def generate_pdf_samples():
    # 1. text.pdf
    doc_text = fitz.open()
    page1 = doc_text.new_page(width=595, height=842) # A4
    page1.insert_text(fitz.Point(50, 80), "LocalPDF 纯文本验证文档 (text.pdf)", fontsize=18)
    page1.insert_text(fitz.Point(50, 120), "第一段：验证段落文本提取、中文字符识别与编码保真度。", fontsize=12)
    page1.insert_text(fitz.Point(50, 150), "第二段：纯本地运行，不依赖云端 API，保证数据不出本地容器。", fontsize=12)
    doc_text.save(str(SAMPLES_DIR / "text.pdf"))
    doc_text.close()

    # 2. table.pdf
    doc_table = fitz.open()
    page_tbl = doc_table.new_page(width=595, height=842)
    page_tbl.insert_text(fitz.Point(50, 60), "员工信息表 (table.pdf)", fontsize=16)
    # 绘制简单表格线
    for y in [100, 130, 160, 190]:
        page_tbl.draw_line(fitz.Point(50, y), fitz.Point(500, y), color=(0.2, 0.2, 0.2), width=1)
    for x in [50, 150, 270, 390, 500]:
        page_tbl.draw_line(fitz.Point(x, 100), fitz.Point(x, 190), color=(0.2, 0.2, 0.2), width=1)

    page_tbl.insert_text(fitz.Point(60, 120), "员工编号", fontsize=11)
    page_tbl.insert_text(fitz.Point(160, 120), "姓名", fontsize=11)
    page_tbl.insert_text(fitz.Point(280, 120), "职位", fontsize=11)
    page_tbl.insert_text(fitz.Point(400, 120), "办公地点", fontsize=11)

    page_tbl.insert_text(fitz.Point(60, 150), "EMP-001", fontsize=10)
    page_tbl.insert_text(fitz.Point(160, 150), "王强", fontsize=10)
    page_tbl.insert_text(fitz.Point(280, 150), "架构师", fontsize=10)
    page_tbl.insert_text(fitz.Point(400, 150), "北京研发中心", fontsize=10)

    page_tbl.insert_text(fitz.Point(60, 180), "EMP-002", fontsize=10)
    page_tbl.insert_text(fitz.Point(160, 180), "赵敏", fontsize=10)
    page_tbl.insert_text(fitz.Point(280, 180), "高级工程师", fontsize=10)
    page_tbl.insert_text(fitz.Point(400, 180), "上海实验室", fontsize=10)

    doc_table.save(str(SAMPLES_DIR / "table.pdf"))
    doc_table.close()

    # 3. scan.pdf (模拟扫描版，将图像直接作为页面底图)
    doc_scan = fitz.open()
    page_scan = doc_scan.new_page(width=595, height=842)
    img_scan = Image.new("RGB", (595, 842), color=(250, 250, 245))
    d_scan = ImageDraw.Draw(img_scan)
    d_scan.text((100, 200), "[SCANNED DOCUMENT] 模拟扫描档案", fill=(80, 80, 80))
    d_scan.text((100, 240), "档案号: ARCH-2024-9988", fill=(80, 80, 80))
    buf = io.BytesIO()
    img_scan.save(buf, "JPEG", quality=80)
    page_scan.insert_image(page_scan.rect, stream=buf.getvalue())
    doc_scan.save(str(SAMPLES_DIR / "scan.pdf"))
    doc_scan.close()

    # 4. invoice.pdf (发票样例)
    doc_inv = fitz.open()
    page_inv = doc_inv.new_page(width=595, height=420) # 横版发票尺寸
    page_inv.insert_text(fitz.Point(220, 50), "电子发票（增值税普通发票）", fontsize=15, color=(0.6, 0.2, 0.2))
    page_inv.insert_text(fitz.Point(40, 90), "发票代码: 011002000111", fontsize=10)
    page_inv.insert_text(fitz.Point(40, 110), "发票号码: 88992211", fontsize=10)
    page_inv.insert_text(fitz.Point(40, 140), "购买方名称: 本地测试科技有限公司", fontsize=10)
    page_inv.insert_text(fitz.Point(40, 160), "纳税人识别号: 91110108MA00000000", fontsize=10)
    page_inv.insert_text(fitz.Point(40, 220), "货物或应税劳务名称: 技术服务费", fontsize=10)
    page_inv.insert_text(fitz.Point(380, 220), "金额: ￥1,000.00", fontsize=10)
    page_inv.insert_text(fitz.Point(40, 280), "价税合计 (大写): 壹仟圆整", fontsize=11, color=(0.2, 0.2, 0.2))
    doc_inv.save(str(SAMPLES_DIR / "invoice.pdf"))
    doc_inv.close()
    print("✓ 已生成 text.pdf, table.pdf, scan.pdf, invoice.pdf")

if __name__ == "__main__":
    print("正在生成 LocalPDF 全套测试样本文件...")
    generate_sample_images()
    generate_docx_samples()
    generate_xlsx_samples()
    generate_pptx_samples()
    generate_pdf_samples()
    print("所有 PRD 第 13 节规定的样本文件已成功生成在:", SAMPLES_DIR)
