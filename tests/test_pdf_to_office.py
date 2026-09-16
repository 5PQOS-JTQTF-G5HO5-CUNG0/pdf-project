import sys
from pathlib import Path
import pytest
import openpyxl
from pptx import Presentation

backend_path = Path(__file__).resolve().parent.parent / "backend"
if str(backend_path) not in sys.path:
    sys.path.insert(0, str(backend_path))

from app.converters.pdf_to_xlsx import PdfToXlsxConverter
from app.converters.pdf_to_pptx import PdfToPptxConverter
from app.converters.router import conversion_router

@pytest.fixture(scope="session")
def setup_samples():
    from tests.generate_samples import SAMPLES_DIR, generate_pdf_samples
    if not (SAMPLES_DIR / "table.pdf").exists():
        generate_pdf_samples()
    return SAMPLES_DIR

def test_router_has_xlsx_and_pptx():
    assert conversion_router.get_converter(".pdf", "xlsx") is not None
    assert conversion_router.get_converter(".pdf", "pptx") is not None

def test_pdf_to_xlsx(setup_samples, tmp_path):
    samples = setup_samples
    src_pdf = samples / "table.pdf"
    out_xlsx = tmp_path / "table_output.xlsx"

    conv = PdfToXlsxConverter()
    res = conv.convert(src_pdf, out_xlsx)
    assert res.exists()
    assert res.stat().st_size > 0

    # 验证 openpyxl 能够正常打开，且工作表中有内容
    wb = openpyxl.load_workbook(str(res))
    sheet = wb.active
    assert sheet is not None
    # 验证提取出了表格文本
    found_any_text = False
    for row in sheet.iter_rows(values_only=True):
        if any(row):
            found_any_text = True
            break
    assert found_any_text is True

def test_pdf_to_pptx_editable(setup_samples, tmp_path):
    samples = setup_samples
    src_pdf = samples / "text.pdf"
    out_pptx = tmp_path / "presentation_editable.pptx"

    conv = PdfToPptxConverter()
    res = conv.convert(src_pdf, out_pptx, options={"mode": "editable"})
    assert res.exists()
    assert res.stat().st_size > 0

    prs = Presentation(str(res))
    assert len(prs.slides) > 0

def test_pdf_to_pptx_presentation(setup_samples, tmp_path):
    samples = setup_samples
    src_pdf = samples / "text.pdf"
    out_pptx = tmp_path / "presentation_hd.pptx"

    conv = PdfToPptxConverter()
    res = conv.convert(src_pdf, out_pptx, options={"mode": "presentation"})
    assert res.exists()
    assert res.stat().st_size > 0

    prs = Presentation(str(res))
    assert len(prs.slides) > 0
