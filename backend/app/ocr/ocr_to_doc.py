import logging
from pathlib import Path
from typing import List
import fitz
from PIL import Image
import docx
from app.ocr.engine import run_ocr_on_image

logger = logging.getLogger("ocr.ocr_to_doc")

def ocr_extract_text(input_path: Path) -> str:
    """提取图片或 PDF 扫描件中的所有文字为纯文本"""
    if not input_path.exists():
        raise FileNotFoundError(f"文件不存在: {input_path}")

    ext = input_path.suffix.lower()
    full_text_lines: List[str] = []

    if ext in (".jpg", ".jpeg", ".png", ".webp"):
        results = run_ocr_on_image(input_path)
        # 按从上到下、从左到右排序
        results.sort(key=lambda r: (min(p[1] for p in r["box"]), min(p[0] for p in r["box"])))
        for item in results:
            full_text_lines.append(item["text"])
    elif ext == ".pdf":
        doc = fitz.open(str(input_path))
        try:
            for page_idx, page in enumerate(doc):
                pix = page.get_pixmap(dpi=200)
                results = run_ocr_on_image(pix.tobytes("png"))
                results.sort(key=lambda r: (min(p[1] for p in r["box"]), min(p[0] for p in r["box"])))
                if len(doc) > 1:
                    full_text_lines.append(f"--- [Page {page_idx + 1}] ---")
                for item in results:
                    full_text_lines.append(item["text"])
        finally:
            doc.close()
    else:
        raise ValueError(f"不支持 OCR 的格式: {ext}")

    return "\n".join(full_text_lines)

def ocr_to_docx(input_path: Path, output_path: Path) -> Path:
    """提取扫描件或图片文字并生成格式整洁的 Word (.docx) 文档"""
    extracted_text = ocr_extract_text(input_path)

    doc = docx.Document()
    doc.add_heading(f"OCR 文字识别结果 - {input_path.name}", level=1)

    paragraphs = extracted_text.split("\n")
    for p in paragraphs:
        p = p.strip()
        if p:
            if p.startswith("--- [Page"):
                doc.add_heading(p, level=2)
            else:
                doc.add_paragraph(p)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    logger.info(f"成功将 OCR 结果保存为 Word: {output_path}")
    return output_path
